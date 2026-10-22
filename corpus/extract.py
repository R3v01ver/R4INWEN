"""把 E:\\小说资料 里的文档全抽成纯文本，落到 corpus/raw/ 下。

- .docx  → python-docx
- .doc   → Word COM 另存为 txt（Office16 在本机）
- .pdf   → pypdf
- .xlsx  → openpyxl
- .epub  → zip 里的 html 去标签
- .rar/.jpg 跳过（rar 记一笔，等他解压）
输出 corpus/_extract_report.txt：每个文件的字数与状态。
"""
import os
import re
import sys
import zipfile

SRC = r"E:\小说资料"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "raw")


def docx_text(p):
    import docx
    d = docx.Document(p)
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def doc_via_word(paths, outdir):
    """一次性开 Word，把一批 .doc 另存成 txt。"""
    import win32com.client as w
    app = w.Dispatch("Word.Application")
    app.Visible = False
    app.DisplayAlerts = False
    done = {}
    try:
        for p in paths:
            name = os.path.splitext(os.path.basename(p))[0]
            dst = os.path.join(outdir, name + ".txt")
            try:
                d = app.Documents.Open(p, ReadOnly=True)
                d.SaveAs2(dst, FileFormat=2)   # 2 = wdFormatText
                d.Close(False)
                done[p] = "ok"
            except Exception as e:
                done[p] = "ERR " + str(e)[:80]
    finally:
        app.Quit()
    return done


def pdf_text(p):
    from pypdf import PdfReader
    r = PdfReader(p)
    return "\n".join((pg.extract_text() or "") for pg in r.pages)


def xlsx_text(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def epub_text(p):
    out = []
    with zipfile.ZipFile(p) as z:
        for n in z.namelist():
            if n.endswith((".html", ".xhtml", ".htm")):
                h = z.read(n).decode("utf-8", "ignore")
                h = re.sub(r"<[^>]+>", "\n", h)
                out.append(h)
    return "\n".join(out)


def main():
    os.makedirs(OUT, exist_ok=True)
    files = []
    for root, _, fs in os.walk(SRC):
        for f in fs:
            files.append(os.path.join(root, f))
    docs, report = [], []
    for p in sorted(files):
        ext = os.path.splitext(p)[1].lower()
        name = os.path.splitext(os.path.basename(p))[0]
        rel = os.path.relpath(p, SRC)
        if ext == ".doc":
            docs.append(p)
            continue
        try:
            if ext == ".docx":
                t = docx_text(p)
            elif ext == ".pdf":
                t = pdf_text(p)
            elif ext == ".xlsx":
                t = xlsx_text(p)
            elif ext == ".epub":
                t = epub_text(p)
            else:
                report.append((rel, 0, "SKIP " + ext))
                continue
            t = re.sub(r"\n{3,}", "\n\n", t).strip()
            open(os.path.join(OUT, name + ".txt"), "w", encoding="utf-8").write(t)
            report.append((rel, len(t), "ok"))
        except Exception as e:
            report.append((rel, 0, "ERR " + str(e)[:90]))

    if docs:
        done = doc_via_word(docs, OUT)
        for p in docs:
            rel = os.path.relpath(p, SRC)
            name = os.path.splitext(os.path.basename(p))[0]
            f = os.path.join(OUT, name + ".txt")
            n = 0
            if os.path.exists(f):
                n = len(open(f, encoding="utf-8", errors="ignore").read())
            report.append((rel, n, done.get(p, "?")))

    report.sort(key=lambda x: -x[1])
    with open(os.path.join(os.path.dirname(OUT), "_extract_report.txt"),
              "w", encoding="utf-8") as fh:
        for rel, n, st in report:
            fh.write(f"{n:>7}  {st:<6}  {rel}\n")
    ok = sum(1 for _, n, s in report if n > 0)
    print(f"文件 {len(report)}，成功 {ok}，总字数 {sum(n for _, n, _ in report)}")


if __name__ == "__main__":
    main()
